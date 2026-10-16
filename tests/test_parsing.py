"""Local parser tests use real PDF/DOCX containers, never parser fakes."""
import hashlib
import io
import zipfile

import pytest
from docx import Document as WordDocument
from pypdf import PdfWriter, get_configuration
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

from campuspulse import parsing
from campuspulse.parsing import DocumentLimitError, DocumentParseError, parse_document


def pdf_bytes(texts, *, encrypt=False, compress=False):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=612, height=792)
        if text is not None:
            font = DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
            })
            stream = DecodedStreamObject()
            escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = stream.flate_encode() if compress else stream
    if encrypt:
        writer.encrypt("test-only-password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes(paragraphs, *, table=False):
    document = WordDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        cells = document.add_table(rows=1, cols=2).rows[0].cells
        cells[0].text = "Provide API documentation."
        cells[1].text = "Ship tests."
        document.add_paragraph("Submit your repository.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def rewrite_docx(original, *, replacements=None, extra=None):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source:
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for entry in source.infolist():
                target.writestr(entry.filename, (replacements or {}).get(entry.filename, source.read(entry)))
            for name, content in (extra or {}).items():
                target.writestr(name, content)
    return output.getvalue()


def test_utf8_paragraphs_preserve_unicode_multiline_and_fingerprint():
    content = "\ufeffҚазақша талап.\r\nSecond line.\r\n\r\nShip tests.".encode("utf-8")
    result = parse_document(content, "C:\\fakepath\\assignment.MD")
    assert result.filename == "assignment.MD"
    assert [(p.number, p.text) for p in result.pages] == [
        (1, "Қазақша талап.\nSecond line."), (2, "Ship tests.")
    ]
    assert result.sha256 == hashlib.sha256(content).hexdigest()


@pytest.mark.parametrize("content, filename, message", [
    (b"", "a.txt", "empty"),
    (b" \n\t\n", "a.txt", "no extractable text"),
    (b"\xff", "a.txt", "UTF-8"),
    (b"abc\x00xyz", "a.txt", "control"),
    (b"abc\x1bxyz", "a.md", "control"),
    (b"text", "a.exe", "Supported"),
    (b"text", "", "filename"),
    (b"text", "bad\x00.txt", "filename"),
    (b"text", "a" * 256 + ".txt", "filename"),
    (b"text", "C:/dir/", "filename"),
    (b"Not a PDF", "a.pdf", "PDF header"),
    (b"%PDF-1.7\nmalformed", "a.pdf", "malformed"),
    (b"not-a-zip", "a.docx", "malformed"),
])
def test_invalid_inputs_fail_clearly(content, filename, message):
    with pytest.raises(DocumentParseError, match=message):
        parse_document(content, filename)


def test_byte_limit_applies_before_parser(monkeypatch):
    monkeypatch.setattr(parsing, "MAX_UPLOAD_BYTES", 10)
    with pytest.raises(DocumentLimitError, match="upload limit"):
        parse_document(b"x" * 11, "a.pdf")


def test_text_is_rejected_instead_of_truncated(monkeypatch):
    monkeypatch.setattr(parsing, "MAX_PAGE_CHARACTERS", 10)
    with pytest.raises(DocumentLimitError, match="page or paragraph"):
        parse_document(b"x" * 11, "a.txt")


def test_total_text_limit(monkeypatch):
    monkeypatch.setattr(parsing, "MAX_TEXT_CHARACTERS", 10)
    with pytest.raises(DocumentLimitError, match="total text"):
        parse_document(b"123456\n\n12345", "a.txt")


def test_paragraph_limit():
    with pytest.raises(DocumentLimitError, match="100 pages or paragraphs"):
        parse_document(b"\n\n".join([b"requirement"] * 101), "a.txt")


def test_real_pdf_keeps_physical_page_numbers_including_blank_pages():
    result = parse_document(pdf_bytes(["Provide tests.", None, "Document the API."]), "a.pdf")
    assert [p.number for p in result.pages] == [1, 2, 3]
    assert result.pages[0].text.strip() == "Provide tests."
    assert result.pages[1].text == ""
    assert result.pages[2].text.strip() == "Document the API."


def test_real_compressed_pdf():
    assert parse_document(pdf_bytes(["Ship tests."], compress=True), "a.pdf").pages[0].text == "Ship tests."


def test_encrypted_pdf_rejected():
    with pytest.raises(DocumentParseError, match="Encrypted PDF"):
        parse_document(pdf_bytes(["Ship tests."], encrypt=True), "a.pdf")


@pytest.mark.parametrize("texts", [[None], [""], ["Ship tests.", ""]])
def test_textless_or_partial_scanned_pdf_rejected(texts):
    with pytest.raises(DocumentParseError, match="extractable text"):
        parse_document(pdf_bytes(texts), "a.pdf")


def test_pdf_page_limit():
    with pytest.raises(DocumentLimitError, match="100 pages"):
        parse_document(pdf_bytes([None] * 101), "a.pdf")


def test_pdf_decompression_limit_and_context_restoration(monkeypatch):
    original_limit = get_configuration().zlib_maximum_output_length
    monkeypatch.setattr(parsing, "MAX_PDF_PAGE_STREAM_BYTES", 200)
    with pytest.raises(DocumentLimitError, match="processing limits"):
        parse_document(pdf_bytes(["A" * 1_000], compress=True), "a.pdf")
    assert get_configuration().zlib_maximum_output_length == original_limit


def test_pdf_total_stream_limit(monkeypatch):
    monkeypatch.setattr(parsing, "MAX_PDF_TOTAL_STREAM_BYTES", 60)
    with pytest.raises(DocumentLimitError, match="processing limit"):
        parse_document(pdf_bytes(["Ship tests.", "Write docs."]), "a.pdf")


def test_image_only_pdf_is_rejected_without_claiming_ocr():
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    image = DecodedStreamObject()
    image.set_data(b"\xff\xff\xff")
    image.update({
        NameObject("/Type"): NameObject("/XObject"),
        NameObject("/Subtype"): NameObject("/Image"),
        NameObject("/Width"): NumberObject(1), NameObject("/Height"): NumberObject(1),
        NameObject("/BitsPerComponent"): NumberObject(8),
        NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): image})
    })
    stream = DecodedStreamObject()
    stream.set_data(b"q 612 0 0 792 0 0 cm /Im0 Do Q")
    page[NameObject("/Contents")] = stream
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(DocumentParseError, match="page 1.*require OCR"):
        parse_document(output.getvalue(), "scan.pdf")


def test_pdf_limits_restore_context_when_text_limit_fails(monkeypatch):
    original = get_configuration()
    monkeypatch.setattr(parsing, "MAX_TEXT_CHARACTERS", 5)
    with pytest.raises(DocumentLimitError, match="text limit"):
        parse_document(pdf_bytes(["Ship tests."]), "a.pdf")
    assert get_configuration() is original


def test_docx_body_and_table_paragraph_references():
    result = parse_document(docx_bytes(["Build a service.", "", "Use Gemini.\nShow evidence."], table=True), "a.docx")
    assert [(p.number, p.text) for p in result.pages] == [
        (1, "Build a service."), (2, ""), (3, "Use Gemini.\nShow evidence."),
        (4, "Provide API documentation."), (5, "Ship tests."), (6, "Submit your repository."),
    ]


def test_docx_empty_rejected():
    with pytest.raises(DocumentParseError, match="no extractable text"):
        parse_document(docx_bytes([]), "a.docx")


def test_docx_paragraph_limit():
    with pytest.raises(DocumentLimitError, match="pages or paragraphs"):
        parse_document(docx_bytes(["Ship tests."] * 101), "a.docx")


@pytest.mark.parametrize("extra_name", ["../outside.xml", "/absolute.xml", "word\\wrong.xml"])
def test_docx_unsafe_archive_names_rejected(extra_name):
    normalized = extra_name.replace("\\", "/")
    content = rewrite_docx(docx_bytes(["Ship tests."]), extra={normalized: b"bad"})
    content = content.replace(normalized.encode(), extra_name.encode())
    with pytest.raises(DocumentParseError, match="unsafe entry"):
        parse_document(content, "a.docx")


def test_docx_zip_bomb_ratio_rejected():
    content = rewrite_docx(docx_bytes(["Ship tests."]), extra={"bomb.bin": b"0" * 100_000})
    with pytest.raises(DocumentLimitError, match="compression ratio"):
        parse_document(content, "a.docx")


@pytest.mark.parametrize("limit, value, message", [
    ("MAX_ZIP_ENTRIES", 2, "too many entries"),
    ("MAX_ZIP_ENTRY_BYTES", 100, "uncompressed size"),
    ("MAX_ZIP_UNCOMPRESSED_BYTES", 100, "uncompressed size"),
])
def test_docx_archive_limits(monkeypatch, limit, value, message):
    content = docx_bytes(["Ship tests."])
    monkeypatch.setattr(parsing, limit, value)
    with pytest.raises(DocumentLimitError, match=message):
        parse_document(content, "a.docx")


def test_docx_duplicate_entries_rejected():
    with pytest.warns(UserWarning, match="Duplicate name"):
        content = rewrite_docx(docx_bytes(["Ship tests."]), extra={"word/document.xml": b"bad"})
    with pytest.raises(DocumentParseError, match="ambiguous"):
        parse_document(content, "a.docx")


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_docx_xml_entities_rejected_even_in_utf16(encoding):
    xml = '<?xml version="1.0" encoding="' + encoding + '"?><!DOCTYPE x [<!ENTITY evil "data">]><x>&evil;</x>'
    content = rewrite_docx(docx_bytes(["Ship tests."]), replacements={"word/document.xml": xml.encode(encoding)})
    with pytest.raises(DocumentParseError, match="document types and entities"):
        parse_document(content, "a.docx")


def test_docx_xml_structure_limit(monkeypatch):
    monkeypatch.setattr(parsing, "MAX_XML_DEPTH", 3)
    with pytest.raises(DocumentLimitError, match="structural"):
        parse_document(docx_bytes(["Ship tests."]), "a.docx")


def test_docx_disguised_zip_rejected():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("not-docx.txt", "Ship tests.")
    with pytest.raises(DocumentParseError, match="DOCX document package"):
        parse_document(output.getvalue(), "a.docx")
